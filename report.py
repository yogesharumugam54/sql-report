"""Turn a SQL database into a formatted Excel sales report.

Creates a small sample SQLite database (customers, products, orders),
runs SQL queries with JOIN and GROUP BY, and writes the results to an
Excel file with one sheet per query.

Usage:
    python report.py                    # builds sample.db, writes sales_report.xlsx
    python report.py --db my.db --output report.xlsx
"""

import argparse
import os
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

QUERIES = {
    "Sales by Product": """
        SELECT p.name AS product,
               SUM(o.quantity) AS units_sold,
               ROUND(SUM(o.quantity * p.price), 2) AS revenue
        FROM orders o
        JOIN products p ON p.id = o.product_id
        GROUP BY p.name
        ORDER BY revenue DESC
    """,
    "Sales by Month": """
        SELECT substr(o.order_date, 1, 7) AS month,
               COUNT(*) AS orders,
               ROUND(SUM(o.quantity * p.price), 2) AS revenue
        FROM orders o
        JOIN products p ON p.id = o.product_id
        GROUP BY month
        ORDER BY month
    """,
    "Top Customers": """
        SELECT c.name AS customer,
               c.city,
               COUNT(*) AS orders,
               ROUND(SUM(o.quantity * p.price), 2) AS total_spent
        FROM orders o
        JOIN customers c ON c.id = o.customer_id
        JOIN products p ON p.id = o.product_id
        GROUP BY c.id
        ORDER BY total_spent DESC
        LIMIT 10
    """,
}


def create_sample_db(path):
    """Build a small database so the project runs without any setup."""
    if os.path.exists(path):
        os.remove(path)
    conn = sqlite3.connect(path)
    conn.executescript("""
        CREATE TABLE customers (id INTEGER PRIMARY KEY, name TEXT, city TEXT);
        CREATE TABLE products  (id INTEGER PRIMARY KEY, name TEXT, price REAL);
        CREATE TABLE orders    (id INTEGER PRIMARY KEY, customer_id INTEGER,
                                product_id INTEGER, quantity INTEGER, order_date TEXT);

        INSERT INTO customers VALUES
            (1,'Asha Rao','Chennai'), (2,'Lars Berg','Växjö'),
            (3,'Priya Raman','Mumbai'), (4,'Sara Lindqvist','Stockholm');
        INSERT INTO products VALUES
            (1,'Keyboard',25.0), (2,'Mouse',12.5), (3,'Monitor',150.0), (4,'Webcam',40.0);
        INSERT INTO orders VALUES
            (1,1,1,2,'2026-01-05'), (2,2,3,1,'2026-01-12'), (3,3,2,4,'2026-02-03'),
            (4,1,3,2,'2026-02-17'), (5,4,4,1,'2026-02-20'), (6,2,1,3,'2026-03-02'),
            (7,3,3,1,'2026-03-15'), (8,4,2,5,'2026-03-28'), (9,1,4,2,'2026-03-30');
    """)
    conn.commit()
    conn.close()


def run_query(conn, sql):
    """Return (column_names, rows) for a SELECT query."""
    cursor = conn.execute(sql)
    headers = [col[0] for col in cursor.description]
    return headers, cursor.fetchall()


def write_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for row in rows:
        sheet.append(row)

    # money columns get a number format, columns get a sensible width
    for index, header in enumerate(headers, start=1):
        letter = get_column_letter(index)
        if header in ("revenue", "total_spent"):
            for cell in sheet[letter][1:]:
                cell.number_format = '#,##0.00'
        longest = max(len(str(c.value)) for c in sheet[letter])
        sheet.column_dimensions[letter].width = longest + 4
    sheet.freeze_panes = "A2"


def build_report(db_path, output_path):
    conn = sqlite3.connect(db_path)
    workbook = Workbook()
    workbook.remove(workbook.active)  # drop the default empty sheet

    for title, sql in QUERIES.items():
        headers, rows = run_query(conn, sql)
        write_sheet(workbook, title, headers, rows)
        print(f"{title}: {len(rows)} rows")

    conn.close()
    workbook.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="SQL to Excel report")
    parser.add_argument("--db", default="sample.db", help="SQLite database file")
    parser.add_argument("--output", default="sales_report.xlsx", help="Excel file to write")
    args = parser.parse_args()

    if args.db == "sample.db":
        create_sample_db(args.db)
    build_report(args.db, args.output)
    print(f"Saved {args.output}")


if __name__ == "__main__":
    main()
